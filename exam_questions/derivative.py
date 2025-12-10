from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Create a new presentation
prs = Presentation()

# -------------------------------------------------
# Slide 1: LSTM Architecture Overview
# -------------------------------------------------
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

# Background for slide 1
bg1 = slide1.shapes.add_shape(
    1,  # auto shape type: rectangle
    Inches(0), Inches(0),
    prs.slide_width, prs.slide_height
)
bg1.fill.solid()
bg1.fill.fore_color.rgb = RGBColor(30, 50, 120)  # Dark blue theme

# Title for slide 1
title_box1 = slide1.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(9), Inches(1.2))
title_frame1 = title_box1.text_frame
title_frame1.text = "LSTM Architecture Overview"
title_p1 = title_frame1.paragraphs[0]
title_p1.font.size = Pt(40)
title_p1.font.bold = True
title_p1.font.color.rgb = RGBColor(255, 255, 255)

# Content for slide 1
content_box1 = slide1.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(10), Inches(4.5))
content_frame1 = content_box1.text_frame

lines1 = [
    "• Input Features: 5",
    "• Hidden Dimension: 94",
    "• Number of LSTM Layers: 3",
    "",
    "Key Highlights:",
    "✓ Multi-layer stacked LSTM for deep temporal learning",
    "✓ 94 memory units capture complex sequential dependencies",
    "✓ Suitable for time-series forecasting / sequence classification"
]

for i, line in enumerate(lines1):
    if i == 0:
        p = content_frame1.paragraphs[0]
    else:
        p = content_frame1.add_paragraph()
    p.text = line
    p.font.size = Pt(24)
    p.font.color.rgb = RGBColor(240, 240, 240)

# -------------------------------------------------
# Slide 2: Project Progress Overview
# -------------------------------------------------
slide2 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

# Background for slide 2
bg2 = slide2.shapes.add_shape(
    1,  # rectangle
    Inches(0), Inches(0),
    prs.slide_width, prs.slide_height
)
bg2.fill.solid()
bg2.fill.fore_color.rgb = RGBColor(20, 40, 100)  # Slightly darker blue

# Title for slide 2
title_box2 = slide2.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(9), Inches(1.2))
title_frame2 = title_box2.text_frame
title_frame2.text = "Project Progress Overview"
title_p2 = title_frame2.paragraphs[0]
title_p2.font.size = Pt(38)
title_p2.font.bold = True
title_p2.font.color.rgb = RGBColor(255, 255, 255)

# Content for slide 2
content_box2 = slide2.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(10), Inches(4.8))
content_frame2 = content_box2.text_frame

lines2 = [
    "✓ Implementation of Mini LSTM Model",
    "   • 5 input features and 94 hidden units",
    "   • 3-layer stacked LSTM architecture",
    "   • Initial training and validation completed",
    "",
    "✓ Simulation of Embedded EIS Measurement System",
    "   • Developed simulation framework for EIS signals",
    "   • Integrated LSTM model with embedded pipeline (concept/simulation stage)",
    "   • Evaluating signal response, accuracy and latency"
]

for i, line in enumerate(lines2):
    if i == 0:
        p = content_frame2.paragraphs[0]
    else:
        p = content_frame2.add_paragraph()
    p.text = line
    p.font.size = Pt(24)
    p.font.color.rgb = RGBColor(240, 240, 240)

# -------------------------------------------------
# Save presentation
# -------------------------------------------------
output_file = "LSTM_EIS_Project_Progress.pptx"
prs.save(output_file)
print(f"Presentation saved as: {output_file}")
